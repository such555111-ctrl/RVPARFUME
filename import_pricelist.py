#!/usr/bin/env python3
"""Импорт прайс-листа «Роза Ветров» (xlsx) в products.json.

Запуск (нужен openpyxl: pip install openpyxl):
    python tools/import_pricelist.py "прайс.xlsx" [--markup 1.0] [--out products.json]

Структура прайса: название бренда — строка в столбце B, под ней позиции
(столбец D — название, F — цена в $). Курс берётся из ячейки E3 прайса.
--markup умножает цену из прайса (1.0 — цена как в прайсе, 1.3 — наценка 30%).
"""
import argparse
import json
import re
import sys

import openpyxl

KEEP_UPPER = {"YSL", "BDK", "MFK", "DKNY", "CK", "EDP", "EDT", "EDC", "XO", "II", "III", "IV", "VI", "NO", "D&G", "MCM", "JHAG"}
CATEGORY = {"M": "Мужская парфюмерия", "L": "Женская парфюмерия", "W": "Женская парфюмерия", "U": "Унисекс"}


def smart_title(text):
    def one(w):
        if w.upper() in KEEP_UPPER or any(ch.isdigit() for ch in w):
            return w.upper() if w.upper() in KEEP_UPPER else w
        return "-".join(p.capitalize() for p in w.split("-"))
    return " ".join(one(w) for w in text.split(" "))


def fmt_vol(v):
    return v.replace(".", ",")


def clean_item(raw):
    s = str(raw).strip().replace("\xa0", " ")
    nospray = bool(re.search(r"non\\?\s*spray|не\s*спрей", s, re.I))
    s = re.sub(r"\(?\s*(non\\?\s*spray|не\s*спрей)\s*\)?", " ", s, flags=re.I)

    gender = None
    g = re.search(r"\((U|L|M|W)\)", s, re.I)
    if g:
        gender = g.group(1).upper()
        s = s.replace(g.group(0), " ")

    size = None
    m = re.search(r"(\d+)\s*\*\s*(\d+(?:[.,]\d+)?)\s*ml", s, re.I)
    if m:  # наборы: 10*2ml, 4*2ML
        size = f"{m.group(1)} × {fmt_vol(m.group(2))} мл"
        s = s.replace(m.group(0), " ")
    else:
        ms = list(re.finditer(r"(\d+(?:[.,]\d+)?)\s*(?:ml|mk|m)\b\.?", s, re.I))
        if ms:
            m = ms[-1]
        else:
            m = re.search(r"(\d+[.,]\d+|\d+)\s*$", s)
        if m:
            size = f"{fmt_vol(m.group(1))} мл"
            s = s[:m.start()] + " " + s[m.end():]

    s = re.sub(r"(?<=[A-Za-z])(EDP|EDT)\b", r" \1", s, flags=re.I)   # MONKEYEDP -> MONKEY EDP
    s = re.sub(r"\s+", " ", s).strip(" .,")
    if not re.search(r"[a-zа-я]", s):
        s = smart_title(s)
    s = re.sub(r"\b(edp|edt|edc)\b", lambda x: x.group(1).upper(), s, flags=re.I)
    return s, gender, size, nospray


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--markup", type=float, default=1.0)
    ap.add_argument("--out", default="products.json")
    args = ap.parse_args()

    ws = openpyxl.load_workbook(args.xlsx).active
    rate = ws["E3"].value
    if not isinstance(rate, (int, float)):
        sys.exit("Не нашёл курс доллара в ячейке E3")

    rows = list(ws.iter_rows(min_row=10, values_only=True))
    entries, brand = [], None
    for r in rows:
        r = list(r) + [None] * 8
        b, d, usd = r[1], r[3], r[5]
        if b and not d:
            brand = re.sub(r"\s+", " ", str(b).strip().replace("`", "'"))
            if brand.isupper():
                brand = smart_title(brand)
            continue
        if d and isinstance(usd, (int, float)):
            entries.append((brand or "", d, usd))

    # Одинаковое название для разных полов (напр. Cerruti 1881 (L) и (M)) —
    # это разные позиции с разной ценой: различаем их пометкой в названии.
    parsed = [(b, clean_item(raw), usd) for b, raw, usd in entries]
    genders_by_key = {}
    for _, (name, gender, size, _), _ in parsed:
        genders_by_key.setdefault((name.lower(), size), set()).add(gender)

    brand_idx = {}
    seen, items = set(), []
    for brand, (name, gender, size, nospray), usd in parsed:
        if len(genders_by_key[(name.lower(), size)]) > 1:
            name += {"M": " (мужской)", "L": " (женский)", "W": " (женский)"}.get(gender, "")
        key = (name.lower(), size)
        if key in seen:          # одна и та же позиция на двух строках прайса
            continue
        seen.add(key)
        brand_idx.setdefault(brand, len(brand_idx))
        desc = f"Бренд: {brand}. " if brand else ""
        desc += f"Пробник {size}." if size else "Пробник."
        if nospray:
            desc += " Без спрея."
        items.append({
            "name": name,
            "brand": brand,
            "category": CATEGORY.get(gender, "Унисекс"),
            "price": int(round(usd * rate * args.markup)),
            "old_price": None,
            "badge": None,
            "sizes": [size] if size else ["Стандарт"],
            "desc": desc,
            "swatch": brand_idx[brand] % 6,
            "images": [],
        })

    # По умолчанию витрина показывает новые (с большим id) первыми, поэтому
    # id идут по убыванию — порядок на витрине совпадает с порядком прайса.
    n = len(items)
    for i, p in enumerate(items):
        p["id"] = n - i
    ordered = [{"id": p["id"], **{k: v for k, v in p.items() if k != "id"}} for p in items]
    with open(args.out, "w", encoding="utf-8") as f:
        f.write("[\n" + ",\n".join(json.dumps(p, ensure_ascii=False) for p in ordered) + "\n]\n")
    print(f"Готово: {n} товаров, {len(brand_idx)} брендов, курс {rate}, наценка x{args.markup} -> {args.out}")


if __name__ == "__main__":
    main()
